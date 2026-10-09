from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from app.errors import AppError


@dataclass
class Passage:
    text: str
    location: dict


def safe_filename(filename: str) -> str:
    name = PurePosixPath(filename.replace("\\", "/")).name
    name = "".join(c for c in name if c.isprintable()).strip()
    if not name or len(name) > 255:
        raise AppError(422, "invalid_filename", "Provide a filename of at most 255 characters")
    return name


def extract(filename: str, data: bytes, max_chars: int) -> list[Passage]:
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt"}:
        raise AppError(415, "unsupported_file", "Supported formats are PDF, DOCX, and TXT")
    if not data:
        raise AppError(422, "empty_file", "The uploaded file is empty")
    passages = []
    total = 0

    def add(text: str, location: dict):
        nonlocal total
        text = text.replace("\x00", "").strip()
        total += len(text)
        if total > max_chars:
            raise AppError(413, "text_limit", "Extracted text exceeds the configured limit")
        if text:
            passages.append(Passage(text, location))

    try:
        if suffix == ".txt":
            encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            decoded = data.decode(encoding)
            if "\x00" in decoded:
                raise AppError(422, "invalid_text", "TXT must contain Unicode text, not binary data")
            lines = decoded.splitlines()
            start = 0
            buffer = []
            for number, line in enumerate(lines, 1):
                if line.strip():
                    if not buffer:
                        start = number
                    buffer.append(line)
                elif buffer:
                    add("\n".join(buffer), {"line_start": start, "line_end": number - 1})
                    buffer = []
            if buffer:
                add("\n".join(buffer), {"line_start": start, "line_end": len(lines)})
        elif suffix == ".pdf":
            if not data.lstrip().startswith(b"%PDF-"):
                raise AppError(422, "invalid_pdf", "File does not have a PDF signature")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise AppError(422, "encrypted_pdf", "Password-protected PDFs are not supported")
            if len(reader.pages) > 2000:
                raise AppError(413, "page_limit", "PDF exceeds 2000 pages")
            for number, page in enumerate(reader.pages, 1):
                add(page.extract_text() or "", {"page": number})
        else:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise AppError(413, "archive_limit", "DOCX expanded content exceeds the limit")
                if "word/document.xml" not in archive.namelist():
                    raise BadZipFile("Not a DOCX")
            document = Document(BytesIO(data))
            for number, paragraph in enumerate(document.paragraphs, 1):
                add(paragraph.text, {"paragraph": number})
            for number, table in enumerate(document.tables, 1):
                for row_number, row in enumerate(table.rows, 1):
                    add(" | ".join(cell.text for cell in row.cells), {"table": number, "row": row_number})
    except AppError:
        raise
    except UnicodeError as exc:
        raise AppError(422, "invalid_encoding", "TXT must be UTF-8 or BOM-marked UTF-16") from exc
    except Exception as exc:
        raise AppError(422, "extraction_failed", "Unable to parse this document; check its format and integrity") from exc
    if not passages:
        raise AppError(422, "no_text", "No extractable text found. Scanned PDFs require OCR before upload.")
    return passages
