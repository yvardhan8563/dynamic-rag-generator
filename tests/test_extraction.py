from io import BytesIO
import re

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.errors import AppError
from app.ingestion.chunking import chunk_passages
from app.ingestion.extractors import Passage, extract, safe_filename


class WordTokenizer:
    def __call__(self, text, **kwargs):
        return {"offset_mapping": [(m.start(), m.end()) for m in re.finditer(r"\S+", text)]}


def pdf_bytes(text="Refunds are available within thirty days."):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_all_formats_and_locations():
    assert extract("a.txt", b"first\nline\n\nsecond", 100)[0].location == {"line_start": 1, "line_end": 2}
    assert extract("a.txt", "Unicode text".encode("utf-16"), 100)[0].text == "Unicode text"
    assert "thirty days" in extract("a.pdf", pdf_bytes(), 1000)[0].text
    assert extract("a.pdf", pdf_bytes(), 1000)[0].location == {"page": 1}
    doc = Document()
    doc.add_paragraph("Paragraph text")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table text"
    output = BytesIO()
    doc.save(output)
    passages = extract("a.docx", output.getvalue(), 1000)
    assert [p.text for p in passages] == ["Paragraph text", "Table text"]
    assert passages[1].location == {"table": 1, "row": 1}


@pytest.mark.parametrize("name,data,code", [
    ("a.exe", b"x", "unsupported_file"), ("a.txt", b"", "empty_file"),
    ("a.txt", b"\xffbad", "invalid_encoding"), ("a.pdf", b"bad", "invalid_pdf"),
    ("a.docx", b"bad", "extraction_failed"), ("a.txt", b" \n ", "no_text"),
    ("a.txt", b"a\x00b", "invalid_text"),
])
def test_invalid_files(name, data, code):
    with pytest.raises(AppError) as error:
        extract(name, data, 100)
    assert error.value.code == code


def test_limits_and_filename():
    with pytest.raises(AppError, match="limit"):
        extract("a.txt", b"a" * 101, 100)
    assert safe_filename("../../private\\policy.txt") == "policy.txt"


def test_chunk_coverage_overlap_and_locations():
    passage = Passage("One two three four five six seven", {"page": 4})
    chunks = chunk_passages([passage], WordTokenizer(), 4, 1, 20)
    assert [c["text"] for c in chunks] == ["One two three four", "four five six seven"]
    assert all(c["location"] == {"page": 4} for c in chunks)
    with pytest.raises(AppError, match="chunk limit"):
        chunk_passages([passage], WordTokenizer(), 4, 1, 1)
